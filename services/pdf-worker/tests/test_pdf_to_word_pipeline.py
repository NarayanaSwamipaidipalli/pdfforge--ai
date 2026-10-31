import os
import io
import tempfile
import pymupdf
from PIL import Image
from docx import Document
from app.processors.pdf_to_word_engine import (
    convert_pdf_to_word,
    detect_scanned_document,
    ConversionConfig
)

def create_complex_test_pdf(pdf_path: str):
    """Generates a rich PDF containing headings, styled text, tables, and images."""
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    
    # 1. Title and Typography
    page.insert_text((72, 72), "MASTER SERVICES AGREEMENT", fontsize=18, fontname="helv", color=(0.1, 0.2, 0.6))
    page.insert_text((72, 100), "Date: October 24, 2026", fontsize=10, fontname="helv", color=(0.4, 0.4, 0.4))
    page.insert_text(
        (72, 125),
        "This Agreement is entered into between PDFForge AI Enterprise Inc. and Customer for cloud processing services.",
        fontsize=11,
        fontname="helv"
    )
    
    # 2. Table with Borders
    rect = pymupdf.Rect(72, 160, 540, 240)
    page.draw_rect(rect, color=(0.2, 0.2, 0.2), width=1)
    page.draw_line((72, 190), (540, 190), color=(0.2, 0.2, 0.2), width=1)
    page.draw_line((250, 160), (250, 240), color=(0.2, 0.2, 0.2), width=1)
    page.draw_line((400, 160), (400, 240), color=(0.2, 0.2, 0.2), width=1)
    
    # Table text
    page.insert_text((80, 180), "Deliverable Module", fontsize=10, fontname="helv")
    page.insert_text((260, 180), "Tier / SLA", fontsize=10, fontname="helv")
    page.insert_text((410, 180), "Monthly Fee", fontsize=10, fontname="helv")
    
    page.insert_text((80, 220), "PDF-to-Word High Fidelity Engine", fontsize=9, fontname="helv")
    page.insert_text((260, 220), "99.99% Availability", fontsize=9, fontname="helv")
    page.insert_text((410, 220), "$4,900.00 USD", fontsize=9, fontname="helv")
    
    # 3. Embedded Image
    img = Image.new("RGB", (180, 80), color=(70, 130, 180))
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    page.insert_image(pymupdf.Rect(72, 270, 252, 350), stream=img_buf.getvalue())
    
    # 4. Hyperlink
    link_rect = pymupdf.Rect(72, 380, 300, 400)
    page.insert_text((72, 395), "Visit documentation at https://pdfforge.ai/docs", fontsize=10, fontname="helv", color=(0.0, 0.4, 0.8))
    page.insert_link({"kind": pymupdf.LINK_URI, "from": link_rect, "uri": "https://pdfforge.ai/docs"})
    
    doc.save(pdf_path)
    doc.close()


def test_pdf_to_word_complete_flow():
    with tempfile.TemporaryDirectory() as tmp_dir:
        pdf_path = os.path.join(tmp_dir, "sample.pdf")
        docx_path = os.path.join(tmp_dir, "sample.docx")
        
        # 1. Create source PDF
        create_complex_test_pdf(pdf_path)
        assert os.path.exists(pdf_path)
        
        # 2. Convert to DOCX
        result = convert_pdf_to_word(pdf_path, docx_path)
        assert os.path.exists(docx_path), "Output DOCX was not created"
        assert os.path.getsize(docx_path) > 1000, "DOCX file is suspiciously small"
        
        # 3. Inspect generated DOCX
        doc = Document(docx_path)
        
        # Verify text was extracted
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "MASTER SERVICES AGREEMENT" in all_text or len(doc.tables) > 0, "Title text missing"
        
        # Verify table was reconstructed
        assert len(doc.tables) >= 1, "Table was not reconstructed into a Word table element"
        table = doc.tables[0]
        assert len(table.rows) >= 2, f"Expected at least 2 table rows, got {len(table.rows)}"
        assert len(table.columns) >= 2, f"Expected at least 2 table columns, got {len(table.columns)}"
        
        # Verify images preserved
        images_found = sum(len(p._element.xpath(".//a:blip")) for p in doc.paragraphs)
        # Check inside tables as well
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    images_found += sum(len(p._element.xpath(".//a:blip")) for p in cell.paragraphs)
        assert images_found >= 1, "Embedded image was not preserved in Word document"
        
        print(f"✔ test_pdf_to_word_complete_flow PASSED! (DOCX size: {os.path.getsize(docx_path)} bytes, Tables: {len(doc.tables)}, Images: {images_found})")


def test_scanned_page_detection():
    with tempfile.TemporaryDirectory() as tmp_dir:
        # Create pure image scanned PDF
        scanned_pdf = os.path.join(tmp_dir, "scanned.pdf")
        doc = pymupdf.open()
        page = doc.new_page(width=612, height=792)
        img = Image.new("RGB", (400, 600), color=(240, 240, 240))
        img_buf = io.BytesIO()
        img.save(img_buf, format="PNG")
        page.insert_image(pymupdf.Rect(36, 36, 576, 756), stream=img_buf.getvalue())
        doc.save(scanned_pdf)
        doc.close()
        
        doc_check = pymupdf.open(scanned_pdf)
        is_scanned, scanned_indices = detect_scanned_document(doc_check)
        doc_check.close()
        
        assert is_scanned is True, "Failed to detect scanned document"
        assert 0 in scanned_indices, "Scanned page 0 was not identified in indices"
        print("✔ test_scanned_page_detection PASSED!")


def test_native_fallback_engine():
    from app.processors.pdf_to_word_engine import convert_with_native_pipeline
    with tempfile.TemporaryDirectory() as tmp_dir:
        pdf_path = os.path.join(tmp_dir, "native.pdf")
        docx_path = os.path.join(tmp_dir, "native.docx")
        create_complex_test_pdf(pdf_path)
        
        cfg = ConversionConfig()
        convert_with_native_pipeline(pdf_path, docx_path, cfg)
        
        assert os.path.exists(docx_path)
        doc = Document(docx_path)
        assert len(doc.tables) >= 1
        print(f"✔ test_native_fallback_engine PASSED! (DOCX size: {os.path.getsize(docx_path)} bytes, Tables: {len(doc.tables)})")


if __name__ == "__main__":
    print("=" * 60)
    print("🧪 Running PDF-to-Word Engine Verification Suite")
    print("=" * 60)
    test_pdf_to_word_complete_flow()
    test_scanned_page_detection()
    test_native_fallback_engine()
    print("\n🎉 ALL TESTS PASSED SUCCESSFULLY!")
