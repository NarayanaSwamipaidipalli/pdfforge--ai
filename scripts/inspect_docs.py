from docx import Document
import fitz

print('=== INSPECTING tinkering lab manual DOCX ===')
p1 = '/tmp/pdfforge/7784086c-097f-48e8-bb2a-26533f23260a/result_7784086c-097f-48e8-bb2a-26533f23260a.docx'
d1 = Document(p1)
print('Paragraphs:', len(d1.paragraphs))
print('Tables:', len(d1.tables))
imgs1 = sum(len(p._element.xpath('.//a:blip')) for p in d1.paragraphs)
print('Images in paragraphs:', imgs1)
for i, p in enumerate(d1.paragraphs[:10]):
    print(f'P{i+1}: {repr(p.text[:60])}')

print('\n=== INSPECTING Document.pdf DOCX ===')
p2 = '/tmp/pdfforge/db77bbff-3569-4859-94e1-f4270a08a8f7/result_db77bbff-3569-4859-94e1-f4270a08a8f7.docx'
d2 = Document(p2)
print('Paragraphs:', len(d2.paragraphs))
print('Tables:', len(d2.tables))
imgs2 = sum(len(p._element.xpath('.//a:blip')) for p in d2.paragraphs)
print('Images in paragraphs:', imgs2)
for i, p in enumerate(d2.paragraphs[:10]):
    print(f'P{i+1}: {repr(p.text[:60])}')

print('\n=== INSPECTING SOURCE Document.pdf ===')
pdf2 = fitz.open('/tmp/pdfforge/db77bbff-3569-4859-94e1-f4270a08a8f7/0_a83196f6_Document.pdf')
print('PDF2 Pages:', len(pdf2))
for i in range(min(3, len(pdf2))):
    print(f'Page {i+1} text length:', len(pdf2[i].get_text()))
    print(f'Page {i+1} images count:', len(pdf2[i].get_images()))
    print(f'Page {i+1} first 100 chars:', repr(pdf2[i].get_text()[:100]))
