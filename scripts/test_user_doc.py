import time, os
from docx import Document
from app.processors.pdf_to_word_engine import convert_pdf_to_word

start = time.time()
pdf_p = '/tmp/pdfforge/db77bbff-3569-4859-94e1-f4270a08a8f7/0_a83196f6_Document.pdf'
docx_p = '/tmp/test_real_doc.docx'

print('Converting real user Document.pdf (82 pages)...')
res = convert_pdf_to_word(pdf_p, docx_p)
dur = time.time() - start

print(f'Conversion complete in {dur:.2f}s! File size: {os.path.getsize(docx_p)} bytes')
doc = Document(docx_p)
print('Generated DOCX Paragraphs:', len(doc.paragraphs))
print('Generated DOCX Tables:', len(doc.tables))
imgs = sum(len(p._element.xpath('.//a:blip')) for p in doc.paragraphs)
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            imgs += sum(len(p._element.xpath('.//a:blip')) for p in cell.paragraphs)
print('Generated DOCX Images:', imgs)

for i in range(min(10, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f'P{i+1} [align={p.alignment}]: {repr(p.text[:70])}')
