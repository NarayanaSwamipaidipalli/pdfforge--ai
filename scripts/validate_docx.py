from docx import Document
d = Document('/tmp/pdfforge/5ee3cefe-2ee9-4aa6-869f-6b76331eae7f/result_5ee3cefe-2ee9-4aa6-869f-6b76331eae7f.docx')
print('=== DOCX VALIDATION ===')
print('Paragraphs:', len(d.paragraphs))
print('Tables:', len(d.tables))
for i, t in enumerate(d.tables):
    print(f'Table {i+1}: rows={len(t.rows)}, cols={len(t.columns)}')
    for r in t.rows:
        row_txt = [c.text.strip() for c in r.cells]
        print('  Row:', row_txt)
for i, p in enumerate(d.paragraphs):
    if p.text.strip():
        print(f'P{i+1} [alignment={p.alignment}]: {repr(p.text.strip())}')
        for r in p.runs:
            print(f'   Run: {repr(r.text)} font={r.font.name} size={r.font.size} bold={r.font.bold}')
